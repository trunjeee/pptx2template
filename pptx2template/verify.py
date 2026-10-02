"""Checks run after every build: structure round-trip and dropped-text diff."""
from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path
from typing import Dict, List, Optional

from .build import LayoutPlan
from .ooxml import RT_SLIDE, Package, qn
from .package import as_presentation
from .parser import Deck, has_real_text, paragraphs_of

ROLE_TO_PPTX = {
    "title": "TITLE", "ctrTitle": "CENTER_TITLE", "subTitle": "SUBTITLE", "body": "BODY", "obj": "OBJECT",
    "pic": "PICTURE", "dt": "DATE", "ftr": "FOOTER", "sldNum": "SLIDE_NUMBER", "chart": "CHART",
    "tbl": "TABLE", "dgm": "ORG_CHART", "media": "MEDIA_CLIP", "clipArt": "CLIP_ART",
}


def deck_texts(deck: Deck) -> List[Counter]:
    """Per slide: multiset of non-empty paragraph strings (taken before the build mutates anything)."""
    out = []
    for slide in deck.slides:
        c = Counter(p for p in paragraphs_of(slide.root.find(qn("p:cSld"))) if has_real_text([p]))
        out.append(c)
    return out


def check_structure(data: bytes, plans: List[LayoutPlan], expect_slides: Optional[Dict[int, str]] = None) -> List[str]:
    """Reopen with python-pptx and compare layouts/placeholders with what was planned."""
    from pptx import Presentation

    problems = []
    prs = Presentation(io.BytesIO(as_presentation(data)))
    masters = list(prs.slide_masters)
    if len(masters) != 1:
        problems.append("expected 1 slide master, found %d" % len(masters))
    layouts = list(masters[0].slide_layouts)
    if len(layouts) != len(plans):
        problems.append("expected %d layouts, found %d" % (len(plans), len(layouts)))
    for layout, plan in zip(layouts, plans):
        if layout.name != plan.name:
            problems.append("layout %r: name is %r" % (plan.name, layout.name))
        got = sorted((ph.placeholder_format.idx, ph.placeholder_format.type.name if ph.placeholder_format.type else "")
                     for ph in layout.placeholders)
        want = sorted((0 if role in ("title", "ctrTitle") else idx, ROLE_TO_PPTX.get(role, ""))
                      for role, idx in plan.placeholders)
        if got != want:
            problems.append("layout %r: placeholders %s, expected %s" % (plan.name, got, want))
    for n, slide in enumerate(prs.slides, start=1):
        layout = slide.slide_layout
        if expect_slides and expect_slides.get(n) and layout.name != expect_slides[n]:
            problems.append("slide %d uses layout %r, expected %r" % (n, layout.name, expect_slides[n]))
        idxs = {ph.placeholder_format.idx for ph in layout.placeholders}
        for ph in slide.placeholders:
            if ph.placeholder_format.idx not in idxs:
                problems.append("slide %d: placeholder %r (idx %d) has no match in layout %r"
                                % (n, ph.name, ph.placeholder_format.idx, layout.name))
    return problems


def check_text(original: List[Counter], data: bytes) -> List[str]:
    """Every paragraph of text in the original deck must still be on its slide."""
    pkg = Package.open(io.BytesIO(data))
    pres_part = pkg.main_part()
    prels = pkg.rels(pres_part)
    lst = pkg.xml(pres_part).find(qn("p:sldIdLst"))
    slides = []
    for el in (lst if lst is not None else []):
        info = prels.get(el.get(qn("r:id")))
        if info and info[0] == RT_SLIDE:
            slides.append(info[1])
    problems = []
    if len(slides) != len(original):
        return ["expected %d slides, found %d" % (len(original), len(slides))]
    for n, (part, want) in enumerate(zip(slides, original), start=1):
        root = pkg.xml(part)
        got = Counter(p for p in paragraphs_of(root.find(qn("p:cSld"))) if p.strip())
        for el in root.iter(qn("p:cNvPr")):
            if el.get("descr"):
                got[el.get("descr")] += 1
        missing = want - got
        for text, count in missing.items():
            problems.append("slide %d: text dropped: %r" % (n, text[:80]))
    return problems


# --------------------------------------------------------------------------- optional visual check
def find_soffice() -> Optional[str]:
    for name in ("soffice", "soffice.exe", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(cand).exists():
            return cand
    return None


def render_pngs(pptx_path: Path, outdir: Path, soffice: str) -> List[Path]:
    """Render via LibreOffice (pdf -> one png per page needs pdftoppm; png export gives first slide only)."""
    outdir.mkdir(parents=True, exist_ok=True)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(pptx_path)],
                   check=True, capture_output=True, timeout=300)
    pdf = outdir / (pptx_path.stem + ".pdf")
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        return []
    subprocess.run([pdftoppm, "-png", "-r", "40", str(pdf), str(outdir / pptx_path.stem)], check=True, timeout=300)
    return sorted(outdir.glob(pptx_path.stem + "*.png"))


def visual_diff(original: Path, data: bytes, threshold: float = 0.02) -> List[str]:
    """Pixel diff of LibreOffice renders. Returns warnings; never raises when tools are missing."""
    soffice = find_soffice()
    if not soffice:
        return ["visual check skipped: LibreOffice (soffice) not found"]
    try:
        from PIL import Image, ImageChops  # optional dependency
    except ImportError:
        return ["visual check skipped: Pillow not installed"]
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        out = tmp / "converted.pptx"
        out.write_bytes(as_presentation(data))
        src = tmp / "original.pptx"
        shutil.copy(original, src)
        try:
            a = render_pngs(src, tmp / "a", soffice)
            b = render_pngs(out, tmp / "b", soffice)
        except Exception as exc:
            return ["visual check skipped: render failed (%s)" % exc]
        if not a or not b:
            return ["visual check skipped: pdftoppm not found"]
        warnings = []
        for n, (pa, pb) in enumerate(zip(a, b), start=1):
            ia, ib = Image.open(pa).convert("L"), Image.open(pb).convert("L")
            if ia.size != ib.size:
                ib = ib.resize(ia.size)
            diff = ImageChops.difference(ia, ib)
            changed = sum(1 for v in diff.getdata() if v > 32) / float(ia.size[0] * ia.size[1])
            if changed > threshold:
                warnings.append("slide %d: %.1f%% of pixels differ from the original render" % (n, changed * 100))
        return warnings
