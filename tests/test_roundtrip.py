import io
import zipfile

import pytest
from lxml import etree
from pptx import Presentation

from pptx2template import convert
from pptx2template.cli import main
from pptx2template.ooxml import CT_TEMPLATE_MAIN, NS
from pptx2template.package import as_presentation, check_wellformed

P = "{%s}" % NS["p"]
A = "{%s}" % NS["a"]


def parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


@pytest.fixture(scope="module")
def sample_result(sample):
    return convert(sample)


def test_verification_passes(sample_result, default_deck):
    assert sample_result.problems == []
    assert convert(default_deck).problems == []


def test_template_content_type(sample_result):
    ct = parts(sample_result.data)["[Content_Types].xml"].decode()
    assert CT_TEMPLATE_MAIN in ct
    assert "presentation.main+xml" not in ct


def test_wellformed(sample_result):
    assert check_wellformed(parts(sample_result.data)) == []


def test_layout_structure(sample_result):
    prs = Presentation(io.BytesIO(as_presentation(sample_result.data)))
    layouts = list(prs.slide_masters[0].slide_layouts)
    assert [l.name for l in layouts] == ["Cover", "Destination", "Layout-TitleTableText-03", "Layout-Title-04"]
    detail = layouts[1]
    assert sorted(ph.placeholder_format.idx for ph in detail.placeholders) == [0, 1, 2, 3]
    names = {sh.name for sh in detail.shapes}
    assert {"Side Band", "Brand Mark", "Separator", "Photo"} <= names
    for n in (2, 3, 4):
        slide = prs.slides[n - 1]
        assert slide.slide_layout.name == "Destination"
        assert {sh.name for sh in slide.shapes} == {"Tag", "Title Text", "Body Text", "Photo"}


def test_every_placeholder_has_lststyle(sample_result):
    """The gotcha: new slides only pick up formatting from lstStyle/defRPr, not from the sample run."""
    for name, data in parts(sample_result.data).items():
        if not name.startswith("ppt/slideLayouts/slideLayout") or not name.endswith(".xml"):
            continue
        root = etree.fromstring(data)
        for sp in root.iter(P + "sp"):
            if sp.find(".//" + P + "ph") is None:
                continue
            d = sp.find(P + "txBody/" + A + "lstStyle/" + A + "lvl1pPr/" + A + "defRPr")
            assert d is not None, "%s: placeholder without lstStyle defRPr" % name


def test_pill_keeps_fill_and_size(sample_result):
    root = etree.fromstring(parts(sample_result.data)["ppt/slideLayouts/slideLayout2.xml"])
    pill = [sp for sp in root.iter(P + "sp") if sp.find(".//" + P + "cNvPr").get("name") == "Tag"][0]
    assert pill.find(P + "spPr/" + A + "solidFill") is not None
    d = pill.find(P + "txBody/" + A + "lstStyle/" + A + "lvl1pPr/" + A + "defRPr")
    assert d.get("sz") == "1200"
    # text colour comes from the autoshape's style (lt1), not the default tx1
    assert d.find(A + "solidFill/" + A + "schemeClr").get("val") == "lt1"


def test_photo_in_shape_becomes_cropped_picture_placeholder(sample_result):
    p = parts(sample_result.data)
    layout = etree.fromstring(p["ppt/slideLayouts/slideLayout2.xml"])
    ph_sp = [sp for sp in layout.iter(P + "sp") if sp.find(".//" + P + "cNvPr").get("name") == "Photo"][0]
    ph = ph_sp.find(".//" + P + "ph")
    assert ph.get("type") == "pic"
    sp_pr = ph_sp.find(P + "spPr")
    assert sp_pr.find(A + "prstGeom").get("prst") == "ellipse"   # inserted photos get the same round crop
    assert sp_pr.find(A + "blipFill") is None                       # the old photo is not baked in
    for n in (2, 3, 4):
        slide = etree.fromstring(p["ppt/slides/slide%d.xml" % n])
        pic = [x for x in slide.iter(P + "pic") if x.find(".//" + P + "cNvPr").get("name") == "Photo"][0]
        assert pic.find(".//" + P + "ph").get("idx") == ph.get("idx")
        assert pic.find(P + "blipFill/" + A + "blip") is not None
        assert pic.find(P + "spPr/" + A + "prstGeom").get("prst") == "ellipse"


def test_plain_photo_becomes_picture_placeholder(sample_result):
    prs = Presentation(io.BytesIO(as_presentation(sample_result.data)))
    cover = prs.slide_masters[0].slide_layouts[0]
    assert [ph.placeholder_format.type.name for ph in cover.placeholders].count("PICTURE") == 1
    photo = [sh for sh in prs.slides[0].placeholders if sh.name == "Cover Photo"][0]
    assert photo.image.blob  # the slide still shows its photo


def test_unique_shape_ids_per_part(sample_result):
    for name, data in parts(sample_result.data).items():
        if name.endswith(".xml") and ("/slides/" in name or "/slideLayouts/" in name):
            ids = [el.get("id") for el in etree.fromstring(data).iter(P + "cNvPr")]
            assert len(ids) == len(set(ids)), name


def test_unicode_text_survives(sample_result):
    prs = Presentation(io.BytesIO(as_presentation(sample_result.data)))
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert "Atlas\xa0Trails Field Guide" in texts
    assert any("’" in t for t in texts)


def test_no_slides(sample):
    r = convert(sample, keep_slides=False)
    assert r.problems == []
    names = parts(r.data)
    assert not any(n.startswith("ppt/slides/") for n in names)
    # photos were slide content and go away; only the small badge (layout chrome) keeps its image
    assert len([n for n in names if n.startswith("ppt/media/")]) == 1


def test_cli(tmp_path, sample, capsys):
    out = tmp_path / "out.potx"
    assert main([str(sample), "-o", str(out), "-q"]) == 0
    assert out.exists()
    assert main([str(sample), "--dry-run"]) == 0
    assert "Destination" in capsys.readouterr().out
