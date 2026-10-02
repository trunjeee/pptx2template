from pptx import Presentation

from pptx2template import analyze

import build_fixtures


def test_grouping_and_names(sample):
    a = analyze(sample)
    groups = [[s.index for s in c.slides] for c in a.clusters]
    assert groups == [[1], [2, 3, 4], [5, 6], [7]]
    names = [c.name for c in a.clusters]
    assert names[0] == "Cover"                    # custom slide name
    assert names[1] == "Destination"          # first custom name in the cluster
    assert names[2] == "Layout-TitleTableText-03"  # shape mix fallback
    assert names[3] == "Layout-Title-04"


def test_representative_is_fullest(sample):
    a = analyze(sample)
    for c in a.clusters:
        assert len(c.rep.shapes) == max(len(s.shapes) for s in c.slides)


def test_tolerance(tmp_path):
    prs = Presentation()
    for dx in (0, 40000, 400000):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        build_fixtures.textbox(s, 600000 + dx, 600000, 4000000, 600000, "Heading", size=32)
        build_fixtures.rect(s, 0, 0, 9144000, 200000, (1, 2, 3))
    path = tmp_path / "tol.pptx"
    prs.save(path)
    assert [[s.index for s in c.slides] for c in analyze(path).clusters] == [[1, 2], [3]]
    assert len(analyze(path, tolerance=500000).clusters) == 1


def test_different_chrome_splits_layouts(tmp_path):
    prs = Presentation()
    for color in ((255, 0, 0), (0, 0, 255)):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        build_fixtures.textbox(s, 600000, 600000, 4000000, 600000, "Heading", size=32)
        build_fixtures.rect(s, 0, 0, 9144000, 200000, color)
    path = tmp_path / "chrome.pptx"
    prs.save(path)
    assert len(analyze(path).clusters) == 2


def test_duplicate_names_are_unique(tmp_path):
    prs = Presentation()
    for size in (32, 12):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld").set("name", "Same")
        build_fixtures.textbox(s, 600000, 600000 * (1 if size == 32 else 4), 4000000, 600000, "X", size=size)
    path = tmp_path / "dup.pptx"
    prs.save(path)
    assert [c.name for c in analyze(path).clusters] == ["Same", "Same 2"]
