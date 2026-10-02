from pptx import Presentation
from pptx.util import Emu
from lxml import etree

from pptx2template import Overrides, analyze
from pptx2template.classify import CHROME, MEDIA, PLACEHOLDER, ShapeOverride

import build_fixtures


def shapes_by_name(analysis, slide_no):
    return {s.name: s for s in analysis.deck.slides[slide_no - 1].shapes}


def test_rules_table(sample):
    a = analyze(sample)
    s = shapes_by_name(a, 2)
    assert s["Tag"].verdict == PLACEHOLDER          # rule 1 beats the pill's solid fill
    assert s["Tag"].reason.startswith("rule 1")
    assert s["Photo"].verdict == PLACEHOLDER         # rule 3: picture fill -> picture placeholder
    assert s["Photo"].role == "pic"
    assert s["Side Band"].verdict == CHROME                 # rule 4: flat fill
    assert s["Separator"].verdict == CHROME               # rule 5: a lone glyph is not text
    assert s["Brand Mark"].verdict == CHROME             # rule 6: group without text
    assert shapes_by_name(a, 1)["Cover Photo"].role == "pic"      # rule 2
    assert shapes_by_name(a, 7)["Badge"].verdict == CHROME        # small picture = logo/icon
    assert shapes_by_name(a, 5)["Table 3"].verdict == MEDIA       # tables cannot be placeholders


def test_title_and_reading_order(sample):
    s = shapes_by_name(analyze(sample), 2)
    assert s["Title Text"].role == "title"                # largest font
    assert (s["Photo"].role, s["Photo"].idx) == ("pic", 1)   # round photo, left of the tag
    assert (s["Tag"].role, s["Tag"].idx) == ("body", 2)      # same row, further right
    assert (s["Body Text"].role, s["Body Text"].idx) == ("body", 3)


def test_existing_title_placeholder_wins(default_deck):
    s = analyze(default_deck).deck.slides[0].shapes
    assert s[0].role in ("title", "ctrTitle")
    assert s[1].role == "subTitle"


def test_chrome_over_photo_stays_on_slide(sample):
    overlay = shapes_by_name(analyze(sample), 1)["Overlay"]
    assert overlay.verdict == MEDIA
    assert "stacking" in overlay.reason


def test_overrides(sample):
    ov = Overrides(shapes={"Side Band": ShapeOverride(force="placeholder", type="body"),
                           "Overlay": ShapeOverride(force="chrome"),
                           "Body Text": ShapeOverride(type="title")},
                   layouts={"cluster_3": "Market Slide"})
    a = analyze(sample, ov)
    s = shapes_by_name(a, 2)
    assert s["Side Band"].verdict == PLACEHOLDER and s["Side Band"].role == "body"
    assert s["Body Text"].role == "title" and s["Title Text"].role == "body"
    assert shapes_by_name(a, 1)["Overlay"].verdict == CHROME   # guard does not undo an override
    assert a.clusters[2].name == "Market Slide"


def test_overrides_file(tmp_path, sample):
    p = tmp_path / "deck.overrides.yaml"
    p.write_text('shapes:\n  "Rectangle 13":\n    force: chrome\nlayouts:\n  cluster_3:\n    name: "Destination"\n')
    ov = Overrides.load(p)
    assert ov.shapes["Rectangle 13"].force == "chrome"
    assert ov.layouts == {"cluster_3": "Destination"}
    (tmp_path / "deck.pptx").write_bytes(b"")
    assert Overrides.find_sidecar(tmp_path / "deck.pptx") == p


def test_animated_chrome_stays_on_slide(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = build_fixtures.rect(s, 0, 0, 1000000, 200000, (10, 20, 30), name="Bar")
    build_fixtures.textbox(s, 0, 500000, 4000000, 500000, "Hello", size=24)
    timing = etree.fromstring(
        '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:anim>'
        '<p:cBhvr><p:cTn id="2"/><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:anim>'
        '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>' % bar.shape_id)
    s._element.append(timing)
    path = tmp_path / "anim.pptx"
    prs.save(path)
    bar_shape = shapes_by_name(analyze(path), 1)["Bar"]
    assert bar_shape.verdict == MEDIA and "animated" in bar_shape.reason
