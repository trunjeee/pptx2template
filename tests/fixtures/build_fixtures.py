"""Generate the sample decks used by the tests.

Everything here is invented: "Atlas Trails" is a fictional travel guide, the logo is
two plain shapes and the "photos" are generated gradients.

Run directly to (re)write the decks next to this file:  python tests/fixtures/build_fixtures.py
"""
from __future__ import annotations

import struct
import sys
import tempfile
import zlib
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).parent

FOREST = (0x0B, 0x3D, 0x2E)
TEAL = (0x0E, 0x6E, 0x6E)
VIOLET = (0x6C, 0x4A, 0xB6)
SAND = (0xF2, 0xC1, 0x4E)


def png(top=(40, 120, 90), bottom=(170, 200, 120), w=96, h=64, texture=True) -> bytes:
    """A small PNG standing in for a photo: a vertical gradient with grainy texture.
    Without texture it is a flat graphic, which the converter treats as decoration."""
    rows = []
    seed = 12345
    for y in range(h):
        t = y / (h - 1)
        row = bytearray(b"\x00")
        for x in range(w):
            seed = (seed * 1103515245 + 12345) & 0x7FFFFFFF
            grain = ((seed >> 16) % 81) - 40 if texture else 0
            row += bytes(max(0, min(255, int(a + (b - a) * t) + grain)) for a, b in zip(top, bottom))
        rows.append(bytes(row))
    raw = b"".join(rows)

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _img(tmp: Path, name: str, top, bottom) -> str:
    p = tmp / name
    p.write_bytes(png(top, bottom))
    return str(p)


def textbox(slide, x, y, w, h, text, size=None, bold=False, color=None, name=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        if size:
            r.font.size = Pt(size)
        r.font.bold = bold
        if color:
            r.font.color.rgb = RGBColor(*color)
    return tb


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(*color)
    s.line.fill.background()
    if name:
        s.name = name
    return s


def picture_fill(slide, x, y, w, h, image_path, shape=MSO_SHAPE.OVAL, name="Photo"):
    """A shape whose fill is a picture (rule 3): here a circular photo."""
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    s.name = name
    _, rid = slide.part.get_or_add_image_part(image_path)
    sp_pr = s._element.spPr
    for child in list(sp_pr):
        if child.tag in (qn("a:solidFill"), qn("a:noFill"), qn("a:gradFill")):
            sp_pr.remove(child)
    blip_fill = etree.Element(qn("a:blipFill"))
    blip = etree.SubElement(blip_fill, qn("a:blip"))
    blip.set(qn("r:embed"), rid)
    stretch = etree.SubElement(blip_fill, qn("a:stretch"))
    etree.SubElement(stretch, qn("a:fillRect"))
    sp_pr.find(qn("a:prstGeom")).addnext(blip_fill)  # fill comes right after the geometry
    s.line.fill.background()
    return s


def brand_mark(slide, x, y):
    """Invented logo: a triangle next to a hexagon, grouped (rule 6)."""
    grp = slide.shapes.add_group_shape()
    grp.name = "Brand Mark"
    a = grp.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Emu(x), Emu(y), Emu(320000), Emu(280000))
    a.fill.solid()
    a.fill.fore_color.rgb = RGBColor(*TEAL)
    a.line.fill.background()
    b = grp.shapes.add_shape(MSO_SHAPE.HEXAGON, Emu(x + 360000), Emu(y), Emu(320000), Emu(280000))
    b.fill.solid()
    b.fill.fore_color.rgb = RGBColor(*SAND)
    b.line.fill.background()
    return grp


def gradient_overlay(slide, x, y, w, h):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    s.name = "Overlay"
    s.fill.gradient()
    s.line.fill.background()
    return s


def set_name(slide, name):
    slide._element.find(qn("p:cSld")).set("name", name)


def sample_deck(path: Path, tmp: Path) -> Path:
    """Travel-guide deck built from text boxes and shapes on blank slides."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    hills = _img(tmp, "hills.png", (60, 120, 170), (120, 170, 90))
    coast = _img(tmp, "coast.png", (40, 150, 190), (230, 210, 160))
    woods = _img(tmp, "woods.png", (20, 70, 40), (90, 140, 70))

    # 1: cover
    s = prs.slides.add_slide(blank)
    set_name(s, "Cover")
    s.shapes.add_picture(hills, 0, 0, Emu(12192000), Emu(6858000)).name = "Cover Photo"
    gradient_overlay(s, 0, 3429000, 12192000, 3429000)
    textbox(s, 700000, 4300000, 9000000, 1000000, "Atlas\xa0Trails Field Guide", size=46, bold=True,
            color=(255, 255, 255), name="Cover Title")
    textbox(s, 700000, 5400000, 9000000, 500000, "Hikers’ edition, 2026", size=20,
            color=(255, 255, 255), name="Cover Subtitle")

    # 2-4: destination pages, identical structure: round photo left, text right, band on the right edge
    stops = [("Mountains", "Alpine Ridge", "Six-hour loop\nViewpoint at 2,100\xa0m", hills),
             ("Coast", "Coral Bay", "Cliff path and coves\nBest at low tide", coast),
             ("Forest", "Pine Hollow", "Shaded and flat\nGood for families", woods)]
    for i, (tag, title, body, img) in enumerate(stops):
        s = prs.slides.add_slide(blank)
        if i == 0:
            set_name(s, "Destination")
        rect(s, 11734800, 0, 457200, 6858000, TEAL, name="Side Band")
        picture_fill(s, 600000, 900000, 4800000, 4800000, img)
        pill = rect(s, 6000000, 900000, 1500000, 360000, VIOLET, MSO_SHAPE.ROUNDED_RECTANGLE, name="Tag")
        pill.text_frame.text = tag
        pill.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        textbox(s, 6000000, 1450000, 5300000, 900000, title, size=40, bold=True, name="Title Text")
        textbox(s, 6000000, 2500000, 5300000, 2200000, body, size=16, name="Body Text")
        textbox(s, 6000000, 5000000, 300000, 300000, "✦", size=14, name="Separator")
        brand_mark(s, 10200000, 6200000)

    # 5-6: standard Title and Content placeholders + a table
    tc = prs.slide_layouts[1]
    for title, bullets in (("Trail Conditions", ["Snow above 1,800 m", "Bridges reopened"]),
                           ("Packing List", ["Two litres of water", "Wind jacket"])):
        s = prs.slides.add_slide(tc)
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        body.add_paragraph().text = bullets[1]
        tbl = s.shapes.add_table(2, 2, Emu(7000000), Emu(5200000), Emu(4000000), Emu(800000)).table
        tbl.cell(0, 0).text = "Month"
        tbl.cell(0, 1).text = "Hikers"
        tbl.cell(1, 0).text = "June"
        tbl.cell(1, 1).text = str(len(title) * 10)

    # 7: closing slide, single big text
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 12192000, 6858000, FOREST, name="Background Block")
    textbox(s, 1000000, 2800000, 10000000, 1200000, "Happy trails", size=54, bold=True, color=(255, 255, 255),
            name="Closing")
    brand_mark(s, 5700000, 5600000)
    s.shapes.add_picture(coast, Emu(11300000), Emu(300000), Emu(500000), Emu(500000)).name = "Badge"

    prs.save(path)
    return path


def default_deck(path: Path) -> Path:
    """Only the stock python-pptx layouts with their placeholders."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Annual Report"
    s.placeholders[1].text = "Prepared by the team"
    for t in ("Agenda", "Results"):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        s.placeholders[1].text = "First point"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Title only"
    prs.save(path)
    return path


def build_all(outdir: Path) -> dict:
    outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        return {
            "sample": sample_deck(outdir / "sample_deck.pptx", Path(tmp)),
            "default": default_deck(outdir / "default_layouts.pptx"),
        }


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE
    for k, v in build_all(out).items():
        print(k, v)
